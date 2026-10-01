from pathlib import Path
import json
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
E = ROOT / 'evidence'
s = json.loads((E / 'summary.json').read_text('utf-8'))
records = {r['label']: r for r in s['transactions']}
txs = {name: json.loads((E / f'{name}.json').read_text('utf-8')) for name in ['transfer','transfer-memo','airdrop']}
d = Document()
sec = d.sections[0]
sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    st = d.styles[name]; st.font.name = 'Arial'; st.font.color.rgb = RGBColor(0,0,0)
    for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
d.styles['Normal'].font.size = Pt(10.5)
d.styles['Normal'].paragraph_format.space_after = Pt(6)
d.styles['Title'].font.size = Pt(22)
d.styles['Heading 1'].font.size = Pt(16)
d.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1','Heading 2']:
    d.styles[name].paragraph_format.space_before = Pt(10)
    d.styles[name].paragraph_format.space_after = Pt(5)

def p(text='',style=None): return d.add_paragraph(text,style)
def h(text): return d.add_heading(text,level=2)
def mono(text, size=9):
    para=p(); para.paragraph_format.space_after=Pt(4)
    r=para.add_run(text);r.font.name='Consolas';r.font.size=Pt(size)
    return para
def link(label,url):
    para=p(); node=OxmlElement('w:hyperlink');node.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'164A7B');pr.append(col);run.append(pr)
    text=OxmlElement('w:t');text.text=label;run.append(text);node.append(run);para._p.append(node)
def table(headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,text in zip(t.rows[0].cells,headers):c.text=text
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(3)
                for run in para.runs:run.font.size=Pt(9.5)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EDF4');c._tc.get_or_add_tcPr().append(shade)
                for run in c.paragraphs[0].runs:run.bold=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

p('Практические работы по Solana','Title')
p('Недели 1–3 • ИТ стартап Итоговый проект 1 • 30 сентября 2026 года')
p('Созданы учебные кошельки, получен 1 тестовый SOL и выполнены две транзакции в Devnet. Транзакции исследованы в Solana Explorer и через RPC. Разработана, собрана и проверена консольная программа на Rust.')
d.add_heading('Неделя 1 Первая транзакция',level=1)
h('Кошельки и получение тестового SOL')
p('Созданы две пары ключей Ed25519 с помощью Keypair.generate из @solana/web3.js. Учебные ключи сохранены локально отдельно от материалов для сдачи. Кошелёк A пополнен через официальный faucet на 1 SOL в Devnet.')
p('A — отправитель и плательщик комиссии');mono(s['sender'])
p('B — получатель');mono(s['recipient'])
link('Подтверждение получения 1 SOL в Solana Explorer',records['airdrop']['explorer'])
h('Перевод 0,001 SOL')
r=records['transfer'];tx=txs['transfer']
table(['Параметр','Результат'],[
('Сеть и статус','Devnet; Success; Finalized'),
('Время UTC / UTC+5','30.09.2026 12:19:53 / 17:19:53'),
('Slot',r['slot']),('Сумма','1 000 000 lamports = 0,001 SOL'),
('Комиссия','5 000 lamports = 0,000005 SOL'),('Подписи / инструкции','1 / 1'),
('Программа','System Program'),('A до / после','1 SOL / 0,998995 SOL'),('B до / после','0 SOL / 0,001 SOL')],[2.0,4.7])
p('Signature');mono(r['signature'],8)
link('Открыть выполненный перевод в Solana Explorer',r['explorer'])
p('Баланс A уменьшился на сумму перевода и комиссию: 1 − 0,001 − 0,000005 = 0,998995 SOL. Баланс B вырос на 0,001 SOL. В RPC поле meta.err равно null, а журнал System Program содержит success.')

d.add_page_break()
d.add_heading('Неделя 2 Transaction Detective',level=1)
p('Для анализа взяты две фактически выполненные транзакции Devnet: обычный перевод и перевод с Memo. Адреса A и B приведены в работе за неделю 1. Account keys описывают участвующие аккаунты, instructions задают действия, program ID определяет исполняемую программу.')
h('Транзакция 1 Обычный перевод SOL')
link('Перевод 0,001 SOL • slot 505901460',records['transfer']['explorer'])
table(['Индекс','Аккаунт','Signer','Writable','Роль'],[
('0','A','Да','Да','Отправитель, fee payer'),('1','B','Нет','Да','Получатель'),('2','System Program','Нет','Нет','Программа перевода')],[.45,1.4,.6,.7,3.55])
p('Program ID System Program');mono('11111111111111111111111111111111')
p('Инструкция №1: system.transfer. Параметры: source = A, destination = B, lamports = 1 000 000. Подпись A разрешает списание. B не подписывает получение SOL. Аккаунты A и B writable, поскольку их балансы меняются. Аккаунт программы используется для исполнения кода и не требует подписи.')
p('Логи показывают вызов System Program и успешное завершение. Внутренних инструкций нет: meta.innerInstructions = []. Потреблено 150 compute units.')
h('Транзакция 2 Перевод и Memo')
r=records['transfer-memo']
mono(r['signature'],8)
link('Перевод 0,002 SOL с Memo • slot 505901468',r['explorer'])
table(['Индекс','Аккаунт','Signer','Writable','Роль'],[
('0','A','Да','Да','Отправитель, fee payer, подписант Memo'),('1','B','Нет','Да','Получатель'),('2','System Program','Нет','Нет','Программа перевода'),('3','Memo Program','Нет','Нет','Проверка и запись заметки в логи')],[.45,1.4,.6,.7,3.55])
p('Program ID Memo');mono('MemoSq4gqABAXKb96qnH8TysNcWxMyWCqXgDLGmfcHr')
p('Инструкция №1: system.transfer на 2 000 000 lamports от A к B. Инструкция №2: spl-memo с текстом ниже; A передан как подписант. Memo проверяет UTF-8 и подпись указанного аккаунта, затем выводит заметку в журнал.')
mono('Solana course weeks 1-3: transaction detective',9)
p('Обе инструкции имеют stackHeight = 1 и выполняются на верхнем уровне одной транзакции. Это последовательные инструкции, а не CPI одной программы в другую.')

d.add_page_break()
d.add_heading('Неделя 2 Результаты анализа',level=1)
table(['Показатель','Перевод','Перевод с Memo'],[
('Статус','Success / Finalized','Success / Finalized'),('Количество account keys','3','4'),('Подписей','1','1'),
('Инструкций верхнего уровня','1','2'),('Программ','System Program','System Program, Memo'),('Комиссия в lamports','5 000','5 000'),
('Compute units consumed','150','31 818'),('Внутренние инструкции','Нет','Нет'),('Изменение A в lamports','−1 005 000','−2 005 000'),('Изменение B в lamports','+1 000 000','+2 000 000')],[2.8,1.65,2.25])
h('Проверка балансов второй транзакции')
p('A: 998 995 000 − 2 000 000 − 5 000 = 996 990 000 lamports, то есть 0,99699 SOL. B: 1 000 000 + 2 000 000 = 3 000 000 lamports, то есть 0,003 SOL. Балансы System Program и Memo Program в этой транзакции не изменились.')
h('Фактические сообщения журнала')
for line in txs['transfer-memo']['meta']['logMessages']:
    mono(line,8)
h('Вывод')
p('Одна транзакция может объединять действия разных программ и использовать одну подпись. Добавление Memo увеличило число аккаунтов и потребление вычислительных ресурсов, но комиссия обеих наблюдаемых транзакций оказалась одинаковой. Из этого не следует, что комиссия всегда равна 5 000 lamports: фактическую комиссию нужно смотреть в meta.fee.')
p('Первая инструкция меняет балансы SOL. Вторая фиксирует заметку в логах и проверяет подписанта. При ошибке любой инструкции обычные изменения состояния транзакции откатываются, однако комиссия сети может быть удержана. Account keys включают и пользовательские аккаунты, и аккаунты программ; количество account keys не равно количеству подписантов.')
h('Подтверждения')
p('В evidence сохранены исходные RPC-ответы transfer.json и transfer-memo.json, данные пополнения airdrop.json, история кошелька, финальные статусы и снимки Explorer. Проверка getSignatureStatuses подтвердила finalized и err = null для всех трёх операций.')

d.add_page_break()
d.add_heading('Неделя 3 Консольная программа на Rust',level=1)
p('Проект solana-account-cli — локальная модель аккаунтов и перевода SOL. Команды demo, balance, convert и transfer демонстрируют основы Rust, применимые в Solana-разработке. При каждом запуске Alice получает начальный баланс 1 SOL, Bob — 0. Сетевые переводы выполнены отдельно в неделях 1–2.')
table(['Конструкция','Применение в проекте'],[
('struct','Account хранит аккаунт, Receipt — результат перевода'),('enum и match','Command задаёт варианты команд; match выбирает обработку'),
('Vec и HashMap','Аргументы CLI; коллекция аккаунтов по имени'),('Result, Option, оператор ?','Проверка ввода, наличие аккаунтов и возврат ошибок'),
('Ссылки и &mut','Чтение и изменение коллекции без передачи владения'),('u64 и checked_add / checked_sub','Суммы в lamports; контроль переполнения и недостатка средств')],[2.6,4.1])
p('Десятичная строка SOL преобразуется в u64 без f64, что исключает округление дробных сумм. До изменения балансов проверяются оба аккаунта, положительная сумма, комиссия и переполнение получателя. Фиксированная комиссия 5 000 lamports служит параметром модели.')
h('Запуск')
mono('cargo run -- demo\ncargo run -- convert 0.000000001\ncargo run -- transfer alice bob 0.001\ncargo test',9)
p('Для Windows приложены готовый bin/solana-account-cli.exe и demo.cmd. Исходники находятся в week3-cli/src/main.rs; внешних Rust-зависимостей нет.')
h('Фактический результат запуска')
mono((E / 'rust-cli-output.txt').read_text('utf-8-sig').strip(),8)
h('Проверка качества')
p('Компилятор: rustc 1.98.1. Сборки debug и release успешны. Все 3 unit-теста прошли: точное преобразование и границы u64; успешный перевод с учётом комиссии; сохранение обоих балансов при недостатке средств и переполнении. Полный текст тестов входит в исходный код.')
p('Ограничения модели: аккаунты обозначаются именами; нет криптографических подписей, RPC, rent и сохранения между запусками. Ошибка модели не списывает комиссию. Это CLI для изучения Rust, а не развёрнутая Solana Program; её развёртывание предусмотрено следующей неделей курса.')
h('Источники')
link('Solana Docs — Instructions и модель транзакций','https://solana.com/docs/core/instructions')
link('Solana Program — Memo','https://www.solana-program.com/docs/memo')
link('The Rust Programming Language — консольные приложения','https://doc.rust-lang.org/book/ch12-00-an-io-project.html')

out=ROOT/'Отчет_недели_1_2_3.docx'
d.save(out)
print(out)

readme=f'''# Выполненные задания за недели 1–3

Главный документ: **Отчет_недели_1_2_3.docx**.

1. Созданы два учебных кошелька, получен 1 SOL в Devnet, выполнен перевод 0.001 SOL.
2. Разобраны два успешных перевода: обычный и перевод 0.002 SOL с Memo.
3. Реализован Rust CLI, выполнены debug/release сборки и три unit-теста.

## Публичные подтверждения

- [Получение тестового SOL]({records['airdrop']['explorer']})
- [Первая транзакция]({records['transfer']['explorer']})
- [Перевод с Memo]({records['transfer-memo']['explorer']})

Все три транзакции подтверждены как Finalized без ошибок.

## Состав папок

- `evidence/` — исходные данные RPC, финальные статусы, снимки Explorer, результаты CLI.
- `week3-cli/` — Rust-исходники, Cargo.toml, Cargo.lock, README, готовая Windows-программа и demo.cmd.
- `network/` — скрипты реальных Devnet-переводов и проверки результатов, package.json и package-lock.json.

Открой `week3-cli/demo.cmd` двойным щелчком для демонстрации третьего задания.
Для независимой проверки через RPC установи Node.js, затем в `network` выполни:

```powershell
npm ci --ignore-scripts
node verify.mjs
```

В архиве нет приватных ключей и node_modules. Оригинальные учебные ключи остаются
только на исходном компьютере в `.private/`.
Скрипт devnet.mjs для повторной отправки требует оригинальные ключи; verify.mjs только читает публичные данные.

## Коротко для защиты

- Account — запись состояния по адресу. Программа получает список нужных аккаунтов.
- Instruction — вызов одной программы с аккаунтами и данными.
- Transaction — подписанный набор инструкций. Signature позволяет найти его в Explorer.
- Signer подтверждает действие подписью; writable разрешает изменение аккаунта.
- В переводе A является signer, writable и fee payer; B — writable без подписи.
- Memo добавляет текст в журнал; во второй транзакции System и Memo вызваны последовательно.
- 1 SOL = 1 000 000 000 lamports. u64 и проверяемые операции исключают дробное округление и переполнение.
- Rust CLI — учебная модель в памяти; реальные Devnet-транзакции находятся в evidence.
'''
(ROOT/'README.md').write_text(readme,encoding='utf-8')
